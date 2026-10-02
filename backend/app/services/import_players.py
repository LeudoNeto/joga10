"""Parse player rosters from pasted text or uploaded files.

Supported inputs:
- Plain text / .txt : one player per line in the format ``Nome - Nota``.
- .csv              : columns Nome, Nota and (optional) Posição.
- .xlsx / .xls      : same columns as CSV, first worksheet.

Parsing is tolerant: decimal comma is accepted, notas are clamped to the
group's range and a missing/invalid nota falls back to the middle of the range
(with a note). Rows without a name are flagged as errors so the caller can
skip them.

``plan_import`` compares the parsed rows with the group's players by name
(accent/case-insensitive): same name and same nota -> ignore, same name and a
different nota -> update the nota, unknown name -> create.

``parse_names`` reads the same inputs but keeps only the names (used to select
the players of a draw from a pasted list, e.g. a WhatsApp roll call).
"""

import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from typing import Optional

from .name_match import normalize

# Accepted header aliases (accent/case-insensitive) mapped to our fields.
HEADER_ALIASES = {
    "name": {"nome", "name", "jogador", "player", "atleta"},
    "skill": {"nota", "skill", "habilidade", "rating", "nivel", "overall", "media"},
    "position": {"posicao", "position", "pos", "funcao", "posição"},
}

# "Nome <sep> Nota" — captures a trailing number as the nota.
_TEXT_LINE = re.compile(
    r"^(?P<name>.*?)\s*[-–—:;\t|]\s*(?P<skill>-?\d+(?:[.,]\d+)?)\s*$"
)


@dataclass
class ParsedRow:
    name: str
    skill: float
    position: Optional[str]
    error: Optional[str] = None
    note: Optional[str] = None
    # False when the nota was missing/invalid and ``skill`` is the default:
    # such a row may create a player but never overwrites an existing nota.
    has_skill: bool = True


@dataclass
class SkillRange:
    low: float = 0.0
    high: float = 10.0

    @property
    def default(self) -> float:
        return round((self.low + self.high) / 2, 2)

    def clamp(self, value: float) -> float:
        return max(self.low, min(self.high, value))


def _norm(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = text.encode("ascii", "ignore").decode("ascii")
    return text.strip().lower()


def _parse_skill(raw, rng: SkillRange) -> tuple[float, Optional[str], bool]:
    """Return ``(nota, note, explicit)``; explicit is False for the default."""
    default = rng.default
    if raw is None or str(raw).strip() == "":
        return default, f"sem nota, usando {default:g}", False
    text = str(raw).strip().replace(",", ".")
    try:
        value = float(text)
    except ValueError:
        return default, f"nota inválida '{raw}', usando {default:g}", False
    if value < rng.low:
        return rng.low, f"nota abaixo de {rng.low:g} ajustada para {rng.low:g}", True
    if value > rng.high:
        return rng.high, f"nota acima de {rng.high:g} ajustada para {rng.high:g}", True
    return value, None, True


def _make_row(name: str, skill_raw, position, rng: SkillRange) -> ParsedRow:
    name = (name or "").strip()
    position = (str(position).strip() or None) if position not in (None, "") else None
    skill, note, explicit = _parse_skill(skill_raw, rng)
    error = None if name else "linha sem nome"
    return ParsedRow(
        name=name, skill=skill, position=position, error=error, note=note, has_skill=explicit
    )


# ---------------------------------------------------------------------------
# Text / .txt
# ---------------------------------------------------------------------------
def parse_text(text: str, rng: SkillRange | None = None) -> list[ParsedRow]:
    rng = rng or SkillRange()
    rows: list[ParsedRow] = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        match = _TEXT_LINE.match(stripped)
        if match:
            rows.append(_make_row(match.group("name"), match.group("skill"), None, rng))
        else:
            # No separator/number: treat the line as a name-only entry (default
            # nota). Leading separator junk (e.g. "- ") is stripped so a line
            # with no real name is flagged instead of imported.
            cleaned = re.sub(r"^[\s\-–—:;|]+", "", stripped).strip()
            rows.append(_make_row(cleaned, None, None, rng))
    return rows


# ---------------------------------------------------------------------------
# Tabular formats (csv / xlsx / xls)
# ---------------------------------------------------------------------------
def _table_columns(table: list[list]) -> tuple[dict[str, Optional[int]], list[list]]:
    """Locate the name/skill/position columns; returns them and the data rows."""
    header = [_norm(cell) for cell in table[0]]
    columns: dict[str, Optional[int]] = {"name": None, "skill": None, "position": None}
    for index, cell in enumerate(header):
        for field, aliases in HEADER_ALIASES.items():
            if cell in aliases and columns[field] is None:
                columns[field] = index

    if any(v is not None for v in columns.values()):
        return columns, table[1:]
    # No recognizable header: assume order name, skill, position.
    width = len(table[0])
    columns = {
        "name": 0,
        "skill": 1 if width > 1 else None,
        "position": 2 if width > 2 else None,
    }
    return columns, table


def _clean_table(table: list[list]) -> list[list]:
    # Drop rows that are entirely empty.
    return [
        row
        for row in table
        if any(str(cell).strip() for cell in row if cell is not None)
    ]


def _cell(row, idx):
    if idx is None or idx >= len(row):
        return None
    value = row[idx]
    return None if value is None else str(value).strip()


def _rows_from_table(table: list[list], rng: SkillRange) -> list[ParsedRow]:
    table = _clean_table(table)
    if not table:
        return []
    columns, data = _table_columns(table)
    return [
        _make_row(
            _cell(row, columns["name"]),
            _cell(row, columns["skill"]),
            _cell(row, columns["position"]),
            rng,
        )
        for row in data
    ]


def _csv_table(content: bytes) -> list[list]:
    text = content.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [list(r) for r in reader]


def _xlsx_table(content: bytes) -> list[list]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    table = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    return table


def _xls_table(content: bytes) -> list[list]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    sheet = book.sheet_by_index(0)
    return [sheet.row_values(r) for r in range(sheet.nrows)]


def parse_csv(content: bytes, rng: SkillRange | None = None) -> list[ParsedRow]:
    return _rows_from_table(_csv_table(content), rng or SkillRange())


def parse_xlsx(content: bytes, rng: SkillRange | None = None) -> list[ParsedRow]:
    return _rows_from_table(_xlsx_table(content), rng or SkillRange())


def parse_xls(content: bytes, rng: SkillRange | None = None) -> list[ParsedRow]:
    return _rows_from_table(_xls_table(content), rng or SkillRange())


def _extension(filename: Optional[str]) -> str:
    return filename.rsplit(".", 1)[-1].lower() if filename and "." in filename else ""


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------
def parse_players(
    filename: Optional[str],
    content: Optional[bytes],
    text: Optional[str],
    rng: SkillRange | None = None,
) -> list[ParsedRow]:
    rng = rng or SkillRange()
    if text and text.strip():
        return parse_text(text, rng)

    if content is not None and filename:
        ext = _extension(filename)
        if ext == "csv":
            return parse_csv(content, rng)
        if ext in ("xlsx", "xlsm"):
            return parse_xlsx(content, rng)
        if ext == "xls":
            return parse_xls(content, rng)
        # txt or anything else: treat as plain text.
        return parse_text(content.decode("utf-8-sig", errors="replace"), rng)

    return []


# ---------------------------------------------------------------------------
# Import plan (preview and confirmation share it)
# ---------------------------------------------------------------------------
ACTION_ORDER = {"update": 0, "create": 1, "ignore": 2}


@dataclass
class PlannedRow:
    row: ParsedRow
    action: str  # update | create | ignore
    player_id: Optional[int] = None  # existing player (update / ignore)
    current_skill: Optional[float] = None
    reason: Optional[str] = None


def _cents(value: float) -> int:
    return int(round(value * 100))


def plan_import(rows: list[ParsedRow], players: list) -> list[PlannedRow]:
    """Decide what each row does against the group's ``players``, matched by
    name ignoring accents, case and punctuation: same nota -> ignore, other
    nota -> update it (the oldest player, if the name is repeated in the
    group), unknown name -> create. Repeated names in the list keep only the
    first line. Result ordered: update, create, ignore (input order inside)."""
    existing: dict[str, list] = {}
    for player in sorted(players, key=lambda p: p.id):
        existing.setdefault(normalize(player.name), []).append(player)

    seen: set[str] = set()
    plan: list[PlannedRow] = []
    for row in rows:
        key = normalize(row.name) if row.name else ""
        if row.error or not key:
            plan.append(PlannedRow(row, "ignore", reason=row.error or "linha sem nome"))
            continue
        if key in seen:
            plan.append(PlannedRow(row, "ignore", reason="nome repetido na lista"))
            continue
        seen.add(key)
        matches = existing.get(key, [])
        if not matches:
            plan.append(PlannedRow(row, "create"))
            continue
        if not row.has_skill:
            target = matches[0]
            plan.append(
                PlannedRow(row, "ignore", target.id, target.skill, "já cadastrado (linha sem nota)")
            )
            continue
        same = next((p for p in matches if _cents(p.skill) == _cents(row.skill)), None)
        if same is not None:
            plan.append(
                PlannedRow(row, "ignore", same.id, same.skill, "já cadastrado com a mesma nota")
            )
        else:
            target = matches[0]
            reason = (
                f"{len(matches)} jogadores com este nome: atualiza o mais antigo"
                if len(matches) > 1
                else None
            )
            plan.append(PlannedRow(row, "update", target.id, target.skill, reason))
    plan.sort(key=lambda p: ACTION_ORDER[p.action])
    return plan


# ---------------------------------------------------------------------------
# Names only (selection import)
# ---------------------------------------------------------------------------
# List markers: "1.", "1)", "01 -", "#3", "- ", "* ", "• "...
_LIST_MARKER = re.compile(
    r"^\s*(?:#?\d{1,3}\s*(?:[.)\-–—:º°]+\s*|\s+)|[-–—*•·>]+\s*)"
)
# Trailing nota, as in the roster format "Neto - 4.80".
_TRAILING_NOTA = re.compile(r"\s*[-–—:;\t|]\s*-?\d+(?:[.,]\d+)?\s*$")


def _is_noise(ch: str) -> bool:
    # emojis and other symbols (✅ ⚽ 👍), control/format chars (incl. the
    # zero-width joiner) and emoji variation selectors / keycap marks
    return (
        unicodedata.category(ch)[0] in ("S", "C")
        or "︀" <= ch <= "️"
        or ch == "⃣"
    )


def _clean_name(raw: str) -> str:
    text = unicodedata.normalize("NFC", raw)
    text = "".join(" " if _is_noise(ch) and ch != "\t" else ch for ch in text)
    text = _LIST_MARKER.sub("", text, count=1)
    text = _TRAILING_NOTA.sub("", text)
    return re.sub(r"\s+", " ", text).strip(" -–—:;|.,")


def parse_names(
    filename: Optional[str], content: Optional[bytes], text: Optional[str]
) -> list[str]:
    """Return the cleaned names in input order (blank lines dropped)."""
    if text and text.strip():
        lines = text.splitlines()
    elif content is not None and filename:
        ext = _extension(filename)
        if ext in ("csv", "xlsx", "xlsm", "xls"):
            reader = {"csv": _csv_table, "xls": _xls_table}.get(ext, _xlsx_table)
            table = _clean_table(reader(content))
            if not table:
                return []
            columns, data = _table_columns(table)
            lines = [_cell(row, columns["name"]) or "" for row in data]
        else:
            lines = content.decode("utf-8-sig", errors="replace").splitlines()
    else:
        return []
    names = [_clean_name(line) for line in lines]
    return [n for n in names if n]
